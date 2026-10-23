"""Stage 1 evaluation (Agents plan Phase 7, §5d step S5–S6): golden weeks + finance Q&A, the D7 gates, ADR-052.

Plan:    docs/plans/Agents_plan.md §5d (approved 2026-10-19), §8 / §8a (G1–G6), D7; acceptance rule ADR-052.
Run:     uv run python tools/run.py pipelines/scripts/agents/Stage1_eval.py              (Phases 1–4: 0 model calls)
         $env:EVAL_REAL_MODELS = '1'; uv run python tools/run.py pipelines/scripts/agents/Stage1_eval.py   (+ Phase 5)
         EVAL_REPEATS (default 2, ADR-052) — repeat 1 uses the llm.py reply cache, repeat k a cache folder of its own
         (Data/processed/agents/llm_cache/eval_r<k>/), so every repeat is a fresh set of calls; the daily quota state is
         shared. A run stopped by the quota resumes from the caches the next day.
         EVAL_TRACE=1 (+ LANGSMITH_API_KEY): each pack and question traced in LangSmith (Phase 8; off by default).

Reads:   the frozen inputs in Data/processed/agents/eval/ (golden_weeks.json, golden_weeks_expected.json,
         golden_findings.json, finance_qa_golden.jsonl — hashes asserted); PostgreSQL read-only through the tools;
         Qdrant + Gemini + Jina (document questions, Phase 5 only); the real models (Phase 5 only).
Writes:  Data/processed/agents/eval/stage1_eval_<date>.json (Phase 5: the results file = the next baseline, ADR-052);
         Notebooks/_scratch/stage1_eval/ (git-ignored: the published eval decks of every run, a partial results file
         while a run is incomplete); temp folders for the packs (P7-4, removed). Checkpoints: InMemorySaver (P7-3).
Re-run:  safe (read-only on the database; Phase 5 results file is named by date and never overwritten).

Phases
1. The frozen inputs (hashes) and the people (the Finance manager runs and approves the packs; the Q&A askers).
2. The scorers on hand-made good / bad decks and answers: each metric must catch its failure.
3. The 11 golden weeks with fake models (0 calls): the harness end to end, every hard gate.
4. The finance Q&A with fake answers (0 calls; the two document questions need live retrieval: Phase 5 only).
5. Real models (EVAL_REAL_MODELS=1): packs + questions, EVAL_REPEATS runs, the results file, the D7 verdict.
Final Checkpoint (narrative: pipelines/checkpoints/agents/Stage1_eval.md).
"""
# %% Setup
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import time
from contextlib import nullcontext
from datetime import date
from pathlib import Path

from langchain_core.messages import AIMessage
from langgraph.checkpoint.memory import InMemorySaver
from langgraph.types import Command
from pptx import Presentation

ROOT = next(p for p in [Path.cwd(), *Path.cwd().parents] if (p / "pyproject.toml").exists() and (p / "Data").is_dir())
sys.path.insert(0, str(ROOT))
from pipelines.agents import llm, roles, tracing  # noqa: E402
from pipelines.agents.deck import sample, slots  # noqa: E402
from pipelines.agents.graphs import finance_qa as qa  # noqa: E402
from pipelines.agents.graphs import models as m  # noqa: E402
from pipelines.agents.graphs import weekly_pack as wp  # noqa: E402
from pipelines.agents.tools import context, db  # noqa: E402

sys.stdout.reconfigure(line_buffering=True)          # progress lines reach the run log as they happen (long runs)
EVAL = ROOT / "Data" / "processed" / "agents" / "eval"
SCRATCH = ROOT / "Notebooks" / "_scratch" / "stage1_eval"
# re-frozen 2026-10-22 (plan §5f F3 / F4 / F5): expected evidence bb3b1a3b… → 7909ccfb… (migration 018) → 1f2bad35…
# (migration 019: audit `review`), findings v2 3a5bde86… → v3 (D-P7-1); the Q&A golden set d7d1e94f… → v2 (D-P7-3 /
# D-P7-4, the scorer notes; 42 questions)
FROZEN = {"golden_weeks.json": "ea13800db38c1a96", "golden_weeks_expected.json": "1f2bad3580b27134",
          "golden_findings.json": "dc6e32484944798a", "finance_qa_golden.jsonl": "5cf72ac33e5765d5"}
N_QUESTIONS = 42
PLANTED = 6                                        # 5 invoice errors (table) + the budget miss of test week 36
REAL = os.environ.get("EVAL_REAL_MODELS") == "1"
REPEATS = int(os.environ.get("EVAL_REPEATS", "2"))
# a candidate change (ADR-052): EVAL_TAG names its runs (cache folders llm_cache/eval_<tag>_r<k>/ for every repeat,
# results stage1_eval_<date>_<tag>.json); EVAL_BASELINE = the baseline results file it is compared with
TAG = os.environ.get("EVAL_TAG", "")
BASELINE = os.environ.get("EVAL_BASELINE", "")
TRACE = os.environ.get("EVAL_TRACE") == "1"     # LangSmith for the real runs (Phase 8, P8-4); off by default

# %% Phase 1 — The frozen inputs and the people
# Every input was frozen before any run (§5d): a changed file stops here, so no result is scored against an answer
# key edited after the fact. The packs run as the Finance manager with the lowest id and are approved by the same
# person (approve only: the evaluation measures the agents, not the approver; G5 edit rate comes from Phase 8).
sha = {}
for name, prefix in FROZEN.items():
    sha[name] = hashlib.sha256((EVAL / name).read_bytes()).hexdigest()
    assert sha[name].startswith(prefix), f"{name} changed since it was frozen ({sha[name][:16]} ≠ {prefix})"
weeks = json.loads((EVAL / "golden_weeks.json").read_text(encoding="utf-8"))["weeks"]
expected = json.loads((EVAL / "golden_weeks_expected.json").read_text(encoding="utf-8"))["evidence"]
findings = json.loads((EVAL / "golden_findings.json").read_text(encoding="utf-8"))["weeks"]
golden_qa = [json.loads(x) for x in (EVAL / "finance_qa_golden.jsonl").read_text(encoding="utf-8").splitlines()]
FIN_MGR = context.load_user(db.fetch(
    "SELECT e.employee_id FROM employees e JOIN roles r USING (role_id) JOIN departments d USING (department_id) "
    "WHERE r.role_name = 'manager' AND d.department_name = 'Finance' ORDER BY e.employee_id LIMIT 1")[0]["employee_id"])
assert len(weeks) == 11 and len(expected) == 11 and set(findings) == set(expected) and len(golden_qa) == N_QUESTIONS
print(f"inputs: {len(weeks)} weeks, {sum(len(v) for v in findings.values())} findings, {len(golden_qa)} questions; "
      f"packs run and approved by employee {FIN_MGR.employee_id} ({FIN_MGR.role}, {FIN_MGR.department})")
print(f"✅ Phase 1 validation passed: the four frozen inputs unchanged (hashes), 11 weeks, {N_QUESTIONS} questions")


# %% Scorers (pure functions: the state of a finished run in, the metrics out)
ERROR_WORDS = re.compile(r"duplicat|error|wrong|mismatch|flagged|overbill|underbill|overcharg|undercharg|"
                         r"exception|mispric|incorrect", re.I)
BELOW_WORDS = re.compile(r"below (its |the )?budget|under (its |the )?budget|missed (its |the )?budget|"
                         r"short of (its |the )?budget|budget miss", re.I)
STRONG_WORDS = re.compile(r"\b(surg|soar|plung|jump|slump|collaps|spik|tumbl|crash|skyrocket)\w*", re.I)
STYLE_ADJ = re.compile(r"\b(surg|soar|plung|slump|collaps|skyrocket|tumbl|stellar|dramatic|massive|huge)\w*", re.I)
REDUNDANT_DIR = re.compile(r"\b(up|down|rose|fell|increased|decreased|grew|declined)\s+[+−-]\s?\d", re.I)
NARRATIVE_SLIDES = range(2, 9)          # the Writer's summary (2) and comments (3–8); slide 9 = external context


def _first_candidates() -> dict[str, str]:
    return {r: c[0].name for r, c in roles.ROLES.items()}


def usage_summary(usage: list[dict]) -> dict:
    first = _first_candidates()
    fresh = [u for u in usage if not u.get("cached") and u.get("outcome") == "ok" and not u["model"].startswith("fake")]
    return {"calls": len(usage), "fresh_calls": len(fresh), "cached_calls": sum(bool(u.get("cached")) for u in usage),
            "tokens": sum(u.get("total_tokens") or 0 for u in usage),
            "max_call_tokens": max((u.get("total_tokens") or 0 for u in usage), default=0),
            "fallbacks": sum(1 for u in usage if u.get("outcome") == "ok" and u.get("role") in first
                             and u["model"] != first[u["role"]] and not u["model"].startswith("fake")),
            "failed_calls": sum(1 for u in usage if u.get("outcome") not in ("ok",)),
            "latency_ms": sum(u.get("latency_ms") or 0 for u in usage)}


def _rendered_bullets(draft: dict, ev: dict) -> list[dict]:
    out = []
    for i, b in enumerate(draft.get("bullets") or []):
        text, bound, errors = slots.render_text(b["text"], ev)
        out.append({"i": i, "slide": b["slide"], "raw": b["text"], "text": text, "bound": bound, "errors": errors,
                    "ids": b.get("evidence_ids") or []})
    return out


def _top3(ev: dict) -> list[str]:
    rows = ev["kpi.category"]["rows"]
    by = "billed_wow_change_sdr" if any(r["billed_wow_change_sdr"] is not None for r in rows) else \
        "budget_variance_sdr"
    return [r["vessel_category"] for r in sorted(rows, key=lambda r: (-abs(r[by] or 0), r["vessel_category"]))[:3]]


def pptx_slide_text(path: Path, slide_no: int) -> str:
    deck = Presentation(str(path))
    texts = []
    for shape in deck.slides[slide_no - 1].shapes:
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
        if getattr(shape, "has_table", False) and shape.has_table:
            texts += [c.text for row in shape.table.rows for c in row.cells]
    return "\n".join(texts)


def _slot_hits(bullets, eid: str, fields: list[str] | None = None, row: str | None = None) -> bool:
    for b in bullets:
        for x in b["bound"]:
            if x.slot.eid != eid or (fields and x.slot.field not in fields):
                continue
            if row is None or (x.slot.key or "").split("/")[0] == row:
                return True
    return False


def finding_met(f: dict, narrative: list[dict], context_bullets: list[dict], ev: dict) -> bool:
    text = "\n".join(b["text"] for b in narrative)
    low = text.lower()
    kind = f["kind"]
    if kind == "suez_context":
        return any(ev.get(e, {}).get("tool") == "suez_transits" for b in context_bullets for e in b["ids"])
    if kind == "flat_week":
        billed_sentences = [b["text"] for b in narrative if re.search(r"bill|revenue", b["text"], re.I)]
        return _slot_hits(narrative, f["evidence"], f["field_any"]) and \
            not any(STRONG_WORDS.search(s) for s in billed_sentences)
    if "row_or_field" in f:
        return any((x.lower() in low) if x.startswith("INV-") else _slot_hits(narrative, f["evidence"], [x])
                   for x in f["row_or_field"])
    if "row_any" in f:
        return any(r.lower() in low or _slot_hits(narrative, f["evidence"], None, r) for r in f["row_any"])
    if "row" in f:
        named = f["row"].lower() in low or _slot_hits(narrative, f["evidence"], None, f["row"])
        return named and (not f.get("field_any") or _slot_hits(narrative, f["evidence"], f["field_any"], f["row"]))
    return _slot_hits(narrative, f["evidence"], f.get("field_any"))


def false_alarms(narrative: list[dict], ev: dict) -> tuple[int, int, list[str]]:
    """(false alarms, findings reported, the false sentences): invoices called wrong that the audit did not flag;
    categories called below budget whose variance is ≥ 0 (a category must be named; the week total is not a claim).
    The original of a flagged duplicate, named in the same sentence ("INV-9000292 duplicates INV-0000292"), is not
    called wrong (§5f F8: the confirmation run's only two "false alarms" were such sentences)."""
    flagged = {r["invoice_number"] for r in ev["audit.week"]["rows"]}
    original_of = {r["invoice_number"]: r.get("duplicate_of") for r in ev["audit.week"]["rows"]}
    variance = {r["vessel_category"]: r["budget_variance_sdr"] for r in ev["kpi.category"]["rows"]}
    fa, ok, bad = 0, 0, []
    for b in narrative:
        if ERROR_WORDS.search(b["text"]):
            named = set(re.findall(r"INV-\d{7}", b["text"]))
            originals = {original_of[i] for i in named & flagged if original_of[i]}
            for inv in named - originals:
                if inv in flagged:
                    ok += 1
                else:
                    fa += 1
                    bad.append(b["text"])
        if BELOW_WORDS.search(b["text"]):
            for cat, var in variance.items():
                if cat.lower() in b["text"].lower():
                    if var is not None and var >= 0:
                        fa += 1
                        bad.append(b["text"])
                    else:
                        ok += 1
    return fa, ok, bad


def claims(narrative: list[dict], ev: dict) -> tuple[int, int, list[str]]:
    """G1, independent of the Validator: a direction word before a change value agrees with its sign; categories
    named in a driver sentence are among the week's top three."""
    top3 = {c.lower() for c in _top3(ev)}
    cats = [r["vessel_category"] for r in ev["kpi.category"]["rows"]]
    n, wrong, bad = 0, 0, []
    for b in narrative:
        for x in b["bound"]:
            if x.slot.direction or not slots.is_change(x.slot.field) or not x.value:
                continue
            before = b["raw"][max(0, x.slot.start - 40):x.slot.start].lower()
            up, down = re.search(qa.UP_WORDS, before), re.search(qa.DOWN_WORDS, before)
            if up or down:
                n += 1
                if (up and not down and x.value < 0) or (down and not up and x.value > 0):
                    wrong += 1
                    bad.append(b["raw"])
        if re.search(r"\bdriv", b["text"], re.I):
            for c in cats:
                if c.lower() in b["text"].lower():
                    n += 1
                    if c.lower() not in top3:
                        wrong += 1
                        bad.append(b["text"])
    return n, wrong, bad


def score_pack(key: str, st: dict, pptx: Path | None) -> dict:
    """One golden week's finished run → its metrics (numbers, planted, findings, false alarms, claims, traceability,
    style, cost)."""
    exp, fl = expected[key], findings[key]
    status = st.get("status")
    published = status == "published"
    ev, draft = st.get("evidence") or {}, st.get("draft") or {}
    bullets = _rendered_bullets(draft, ev) if draft else []
    narrative = [b for b in bullets if b["slide"] in NARRATIVE_SLIDES]
    context_b = [b for b in bullets if b["slide"] == 9]
    ev_same = all(ev.get(k) == v for k, v in exp.items())
    n_bound = sum(len(b["bound"]) for b in bullets)
    slot_errors = sum(len(b["errors"]) for b in bullets)
    numbers_ok = published and ev_same and slot_errors == 0 and bool((st.get("verdict") or {}).get("passed"))
    # planted (findings v3, D-P7-1): the 5 invoice errors in the slide-8 table + the budget miss of test week 36,
    # found when a narrative sentence names the category and calls it below budget / a miss
    planted_tab = [f for f in fl if f["kind"] == "invoice_error" and f["where"] == "table"]
    planted_miss = [f for f in fl if f["kind"] == "budget_miss" and f.get("planted")]
    planted = planted_tab + planted_miss
    table = pptx_slide_text(pptx, 8) if (published and pptx and pptx.exists()) else ""
    planted_found = ([f["id"] for f in planted_tab if f["row"] in table]
                     + [f["id"] for f in planted_miss if published and any(
                         BELOW_WORDS.search(b["text"]) and f["row"].lower() in b["text"].lower() for b in narrative)])
    recall_items = [f for f in fl if f["where"] == "narrative" and not f.get("planted")
                    and f["kind"] not in ("no_exception", "no_false_miss")]
    met = [f["id"] for f in recall_items if published and finding_met(f, narrative, context_b, ev)]
    fa, reported, fa_text = false_alarms(narrative, ev) if published else (0, 0, [])
    n_claims, wrong_claims, claim_text = claims(narrative, ev) if published else (0, 0, [])
    traced = sum(1 for b in bullets for x in b["bound"] if x.slot.eid in b["ids"]
                 and (ev[x.slot.eid].get("source") is not None))
    titles = [t["title"] for t in draft.get("slide_titles") or []]
    return {"week": key, "status": status, "stop_reason": st.get("stop_reason"), "rewrites": st.get("rewrites", 0),
            "published": published, "evidence_equal_frozen": ev_same, "numbers_bound": n_bound,
            "slot_errors": slot_errors, "numbers_ok": numbers_ok,
            "planted": [f["id"] for f in planted], "planted_found": planted_found,
            "findings": [f["id"] for f in recall_items], "findings_met": met,
            "false_alarms": fa, "findings_reported": reported, "false_alarm_text": fa_text[:3],
            "claims": n_claims, "claims_wrong": wrong_claims, "claim_text": claim_text[:3],
            "traceable": traced, "style": {
                "redundant_direction": sum(len(REDUNDANT_DIR.findall(b["text"])) for b in bullets),
                "empty_bullets": sum(1 for b in narrative if not re.search(r"\d", b["text"])
                                     and not re.search(r"INV-|[A-Z][a-z]+ (Carriers|Tankers|Ships|Vessels)",
                                                       b["text"])),
                "title_adjectives": sum(1 for t in titles if STYLE_ADJ.search(t))},
            "usage": usage_summary(st.get("llm_usage") or [])}


def _week_in(text: str, week: int) -> bool:
    return bool(re.search(rf"(?<!\d){week}(?!\d)", text.translate(qa.ARABIC_DIGITS)))


def element_met(el: dict, st: dict) -> bool:
    rendered = st.get("rendered") or ""
    if el["kind"] == "value":
        calls = st.get("calls") or {}
        for b in st.get("bound") or []:
            c = calls.get(b["eid"])
            if not c or c["tool"] != el["tool"] or any(c["args"].get(k) != v for k, v in el["args"].items()):
                continue
            if b["field"] in el["field_any"] and (el["row"] is None or (b["row_label"] or "").split("/")[0] == el["row"]):
                return True
        return False
    if el["kind"] in ("name", "direction"):
        return any(t.lower() in rendered.lower() for t in el["text_any"])
    return bool(re.search(el["regex"], rendered, re.I))


def score_qa(item: dict, st: dict) -> dict:
    status = st.get("status") or "error"
    rendered = st.get("rendered") or ""
    bound = st.get("bound") or []
    expected_status = {item["expected_status"], *item.get("status_also", [])}
    out = {"id": item["id"], "lang": item["lang"], "expected": item["expected_status"], "status": status,
           "status_ok": status in expected_status, "problems": st.get("problems") or [], "rendered": rendered[:600],
           "usage": usage_summary(st.get("llm_usage") or [])}
    if item["expected_status"] == "answered":
        met = [element_met(e, st) for e in item["elements"]] if status == "answered" else [False] * len(item["elements"])
        rel = item["relevance"]
        out.update(completeness=sum(met) / len(met) if met else 0.0, elements_met=sum(met),
                   relevance=float(status == "answered"
                                   and (rel["week"] is None or _week_in(rendered, rel["week"]))
                                   and any(w.lower() in rendered.lower() for w in rel["metric_any"])),
                   numbers_ok=status == "answered" and not out["problems"],
                   traceable=all(b["eid"] in ((st.get("answer") or {}).get("evidence_ids") or []) for b in bound))
    else:
        leak = bool(bound) or (bool(item.get("must_not_regex")) and bool(re.search(item["must_not_regex"], rendered)))
        out.update(refusal_ok=out["status_ok"] and not leak, leak=leak)
    return out


def aggregate(packs: list[dict], answers: list[dict]) -> dict:
    pub = [p for p in packs if p["published"]]
    planted = sum(len(p["planted"]) for p in packs)
    found = sum(len(p["planted_found"]) for p in packs)
    rec_n = sum(len(p["findings"]) for p in packs)
    rec_m = sum(len(p["findings_met"]) for p in packs)
    fa, rep = sum(p["false_alarms"] for p in packs), sum(p["findings_reported"] for p in packs)
    cl, cw = sum(p["claims"] for p in packs), sum(p["claims_wrong"] for p in packs)
    nb, tr = sum(p["numbers_bound"] for p in pub), sum(p["traceable"] for p in pub)
    ans = {lang: [a for a in answers if a["lang"] == lang and a["expected"] == "answered"] for lang in ("en", "ar")}
    ref = [a for a in answers if a["expected"] != "answered"]
    said_no = [a for a in answers if a["status"] in ("not_in_data", "not_permitted")]
    seen = {w["week"] for w in weeks if w["seen"]}

    def mean(xs):
        return round(sum(xs) / len(xs), 4) if xs else None
    return {
        "packs_published": f"{len(pub)} / {len(packs)}",
        "deck_numbers": mean([float(p["numbers_ok"]) for p in pub]),
        "planted_found": f"{found} / {planted}",
        "false_alarm_rate": round(fa / (fa + rep), 4) if (fa + rep) else 0.0,
        "claim_accuracy": round(1 - cw / cl, 4) if cl else 1.0,
        "traceability": round(tr / nb, 4) if nb else 1.0,
        "narrative_recall": round(rec_m / rec_n, 4) if rec_n else None,
        "narrative_recall_seen": mean([len(p["findings_met"]) / len(p["findings"]) for p in packs
                                       if p["findings"] and int(p["week"].split(":")[1]) in seen]),
        "narrative_recall_unseen": mean([len(p["findings_met"]) / len(p["findings"]) for p in packs
                                         if p["findings"] and int(p["week"].split(":")[1]) not in seen]),
        "qa_completeness_en": mean([a["completeness"] for a in ans["en"]]),
        "qa_completeness_ar": mean([a["completeness"] for a in ans["ar"]]),
        "qa_relevance_en": mean([a["relevance"] for a in ans["en"]]),
        "qa_relevance_ar": mean([a["relevance"] for a in ans["ar"]]),
        "qa_numbers": mean([float(a["numbers_ok"]) for a in ans["en"] + ans["ar"] if a["status"] == "answered"]),
        "qa_traceability": mean([float(a["traceable"]) for a in ans["en"] + ans["ar"] if a["status"] == "answered"]),
        "abstention_recall": mean([float(a["refusal_ok"]) for a in ref]),
        "abstention_precision": round(sum(a["expected"] != "answered" for a in said_no) / len(said_no), 4)
        if said_no else None,
        "rbac_leakage": sum(a.get("leak", False) for a in ref),
        "style": {k: sum(p["style"][k] for p in pub) for k in ("redundant_direction", "empty_bullets",
                                                               "title_adjectives")},
        "usage": {k: sum(x["usage"][k] for x in packs + answers) for k in ("calls", "fresh_calls", "cached_calls",
                                                                          "tokens", "fallbacks", "failed_calls")},
        "rewrites": sum(p["rewrites"] for p in packs)}


def gates(agg: dict) -> dict:
    """D7 + §8a hard gates (ADR-052)."""
    f, n = (int(x) for x in agg["planted_found"].split(" / "))
    return {"deck_numbers_100": agg["deck_numbers"] == 1.0,
            "planted_all": f == n == PLANTED,          # 6 since findings v3 (D-P7-1); the c3 baseline had 5 / 5
            "false_alarms_le_5pct": agg["false_alarm_rate"] <= 0.05,
            "claims_100": agg["claim_accuracy"] == 1.0,
            "traceability_100": agg["traceability"] == 1.0 and agg["qa_traceability"] in (None, 1.0),
            "rbac_leakage_0": agg["rbac_leakage"] == 0,
            "abstention_recall_100": agg["abstention_recall"] == 1.0}


# %% Phase 2 — The scorers on hand-made decks and answers
# A scorer that never fails proves nothing: every metric is first shown a good case and a case it must catch. The
# decks are the sample slot draft (deck/sample.py) on the frozen evidence, published by hand (the state a real run
# leaves), then changed one way each; the answers are written from the golden elements.
def fake_state(key: str, draft: dict | None = None, **over) -> dict:
    w = int(key.split(":")[1])
    ev = json.loads(json.dumps(expected[key]))
    return {"status": "published", "evidence": ev, "draft": draft or sample.slot_draft(w, ev),
            "verdict": {"passed": True}, "rewrites": 0, "llm_usage": [], **over}


def with_bullet(st: dict, slide: int, text: str, ids: list[str]) -> dict:
    d = json.loads(json.dumps(st["draft"]))
    d["bullets"].append({"slide": slide, "text": text, "evidence_ids": ids})
    return {**st, "draft": d}


good = score_pack("live:47", fake_state("live:47"), None)
cases = {
    "good": (good["numbers_ok"], good["false_alarms"], good["claims_wrong"], good["traceable"] == good["numbers_bound"]),
    "evidence_tampered": score_pack("live:47", {**fake_state("live:47"), "evidence": {
        **expected["live:47"], "kpi.weekly": {**expected["live:47"]["kpi.weekly"], "rows": [
            {**expected["live:47"]["kpi.weekly"]["rows"][0], "billed_sdr": 1.0}]}}}, None)["numbers_ok"],
    "false_invoice": score_pack("test:49", with_bullet(fake_state("test:49"), 8,
                                                       "INV-0000001 was billed twice (a duplicate)", ["audit.week"]),
                                None)["false_alarms"],
    # §5f F8: a flagged duplicate named with its original — one finding reported, no alarm; an unflagged invoice in
    # the same sentence is still one
    "duplicate_with_original": (lambda s, base: (s["false_alarms"], s["findings_reported"] - base))(score_pack(
        "test:23", with_bullet(fake_state("test:23"), 2, "INV-9000292 was flagged as a duplicate of INV-0000292",
                               ["audit.week"]), None),
        score_pack("test:23", fake_state("test:23"), None)["findings_reported"]),
    "duplicate_with_other": score_pack("test:23", with_bullet(
        fake_state("test:23"), 2, "INV-9000292 and INV-0000001 were flagged as duplicates", ["audit.week"]),
        None)["false_alarms"],
    # week 36 of the copy: Special Floating Units is above budget (+25.53 %), LNG is the planted miss (−33.33 %)
    "false_miss": score_pack("test:36", with_bullet(fake_state("test:36"), 5,
                                                    "Special Floating Units billed below budget",
                                                    ["kpi.category"]), None)["false_alarms"],
    "planted_miss_found": "t36-b" in score_pack("test:36", with_bullet(
        fake_state("test:36"), 5, "Liquefied Natural Gas (LNG) Carriers came in below budget", ["kpi.category"]),
        None)["planted_found"],
    "planted_miss_silent": "t36-b" in score_pack("test:36", fake_state("test:36"), None)["planted_found"],
    "wrong_direction": score_pack("live:47", with_bullet(fake_state("live:47"), 3,
                                                         "Billing rose {kpi.weekly:billed_wow_pct|%1} this week",
                                                         ["kpi.weekly"]), None)["claims_wrong"],
    "non_driver": score_pack("live:47", with_bullet(fake_state("live:47"), 4,
                                                    "Revenue was driven by Vehicles Carriers", ["kpi.category"]),
                             None)["claims_wrong"],
    "untraced": (lambda s: s["traceable"] < s["numbers_bound"])(score_pack("live:47", with_bullet(
        fake_state("live:47"), 3, "Overdue {kpi.weekly:overdue_sdr|M}", []), None)),
    "budget_miss_named": "l47-a" in score_pack("live:47", with_bullet(
        fake_state("live:47"), 2, "Billing was {kpi.weekly:budget_variance_pct|%1} against budget", ["kpi.weekly"]),
        None)["findings_met"],
    "flat_week_strong": "l41-a" in score_pack("live:41", with_bullet(
        fake_state("live:41"), 2, "Billing surged {kpi.weekly:billed_wow_pct|%1|dir}", ["kpi.weekly"]),
        None)["findings_met"],
    "not_published": score_pack("live:47", {**fake_state("live:47"), "status": "validation_failed"},
                                None)["findings_met"],
}


def qa_state(item: dict, text: str, status: str = "answered", bound: list | None = None, calls=None) -> dict:
    return {"status": status, "rendered": text, "bound": bound or [], "calls": calls or {}, "problems": [],
            "answer": {"evidence_ids": [b["eid"] for b in bound or []]}, "llm_usage": []}


def good_answer(item: dict) -> dict:
    """A right answer written from the golden elements: one bound slot per value element, the names and words."""
    bound, calls, words = [], {}, []
    for i, el in enumerate(item["elements"]):
        if el["kind"] == "value":
            eid = f"insight.c{i}"
            calls[eid] = {"tool": el["tool"], "args": dict(el["args"])}
            f = el["field_any"][0]
            bound.append({"eid": eid, "row": 0, "row_label": el["row"], "field": f, "value": el["values"][f],
                          "text": slots.format_value(f, el["values"][f])})
        elif el["kind"] in ("name", "direction"):
            words.append(el["text_any"][0])
        else:
            words.append("€940.4 million")
    w = item["relevance"]["week"]
    text = " ".join([*(b["text"] for b in bound), *words, f"week {w}" if w else "",
                     item["relevance"]["metric_any"][0]])
    return qa_state(item, text, bound=bound, calls=calls)


by_id = {i["id"]: i for i in golden_qa}
ans_good = [score_qa(i, good_answer(i)) for i in golden_qa if i["expected_status"] == "answered"]
q01, r01, r03, a01 = by_id["q01"], by_id["r01"], by_id["r03"], by_id["a01"]
qa_cases = {
    "good_complete": min(a["completeness"] for a in ans_good), "good_relevant": min(a["relevance"] for a in ans_good),
    "wrong_week_value": score_qa(q01, {**good_answer(q01), "calls": {"insight.c0": {
        "tool": "finance_kpi", "args": {"report_week": 22, "dataset": "weekly"}}}})["completeness"],
    "refusal_ok": score_qa(r01, qa_state(r01, "Your role has no access to receivables.", "not_permitted"))["refusal_ok"],
    "refusal_with_number": score_qa(r01, qa_state(r01, "AR was 47.63 M SDR.", "answered",
                                                  bound=good_answer(q01)["bound"]))["refusal_ok"],
    "hbr_leak": score_qa(r03, qa_state(r03, "Revenue grew to €940.4 million.", "not_in_data"))["leak"],
    "abstain_ok": score_qa(a01, qa_state(a01, "NileGate's data covers 2025 only.", "not_in_data"))["refusal_ok"],
    "answered_instead": score_qa(a01, qa_state(a01, "Billing was high.", "answered"))["refusal_ok"],
}
actual = {"pack": cases, "qa": qa_cases}
expected_cases = {"pack": {"good": (True, 0, 0, True), "evidence_tampered": False, "false_invoice": 1,
                           "duplicate_with_original": (0, 1), "duplicate_with_other": 1, "false_miss": 1, "planted_miss_found": True, "planted_miss_silent": False,
                           "wrong_direction": 1, "non_driver": 1, "untraced": True,
                           "budget_miss_named": True, "flat_week_strong": False, "not_published": []},
                  "qa": {"good_complete": 1.0, "good_relevant": 1.0, "wrong_week_value": 0.0, "refusal_ok": True,
                         "refusal_with_number": False, "hbr_leak": True, "abstain_ok": True,
                         "answered_instead": False}}
assert actual == expected_cases, actual
# the Q&A checker (c3, D-P7-15): digits inside a tool's own text (a vessel named "… 165", a citation page) are names,
# not typed figures; a NileGate amount typed as text is still refused
od = {"insight.c1": {"tool": "finance_kpi", "status": "ok", "source": {"kind": "view", "name": "finance.kpi_overdue_list",
                                                                       "filter": "f"}, "note": None,
                     "rows": [{"invoice_number": "INV-0000723", "vessel_name": "Liquefied Voyager 165",
                               "amount_sdr": 673471.35, "days_past_due": 49}]}}
name_ok = qa.check_answer({"status": "answered", "evidence_ids": ["insight.c1"],
                           "answer": "INV-0000723 of the vessel Liquefied Voyager 165: "
                                     "{insight.c1[INV-0000723]:amount_sdr}."}, od, "Largest overdue in week 49?")[2]
typed = qa.check_answer({"status": "answered", "evidence_ids": ["insight.c1"],
                         "answer": "INV-0000723 of Liquefied Voyager 165: 673,471.35 SDR."}, od,
                        "Largest overdue in week 49?")[2]
assert name_ok == [] and len(typed) == 1 and "673471.35" in typed[0], (name_ok, typed)
print(f"Phase 2: {actual}")
print("✅ Phase 2 validation passed: every scorer passes its good case and catches its failure (tampered evidence, "
      "false invoice, false budget miss, wrong direction, non-driver, untraced number, strong word on a flat week, "
      "wrong week, refusal with a number, a level-3 leak, an answer instead of an abstention)")


# %% Runners
def eval_trace(thread_id: str, user, label: str):
    """LangSmith only with EVAL_TRACE=1 and a key (Phase 8, §5e P8-4): one evaluation ≈ 400 calls."""
    if not TRACE:
        return nullcontext(False)
    return tracing.trace(thread_id, entry="eval", run_as=user.employee_id, launched_by="eval",
                         extra_tags=(f"eval:{TAG or 'untagged'}", label))


def run_pack(graph, key: str, reports: Path, user=FIN_MGR) -> tuple[dict, Path | None]:
    ds, w = key.split(":")[0], int(key.split(":")[1])
    cfg = {"configurable": {"thread_id": f"weekly_pack:eval:{key}:{time.time_ns()}"}}
    ctx = context.RunContext(user, ds)
    with eval_trace(cfg["configurable"]["thread_id"], user, f"golden:{key}"):
        graph.invoke({"report_week": w}, cfg, context=ctx)
        st = graph.get_state(cfg)
        if st.next and "approval" in st.next and (st.values.get("verdict") or {}).get("passed"):
            graph.invoke(Command(resume={"decision": "approve", "approver_employee_id": user.employee_id}), cfg,
                         context=ctx)
            st = graph.get_state(cfg)
    vals = dict(st.values)
    if vals.get("status") == "awaiting_approval" and not (vals.get("verdict") or {}).get("passed"):
        vals["status"] = "validation_failed"
    pptx = Path(vals["published"]["folder"]) / "deck.pptx" if vals.get("status") == "published" else None
    return vals, pptx


def run_qa(graph, item: dict) -> dict:
    user = context.load_user(item["employee_id"])
    thread = f"qa:{item['id']}:{time.time_ns()}"
    with eval_trace(thread, user, f"qa:{item['id']}"):
        return qa.ask(graph, item["question"], user, item["lang"], thread=thread)


class FakePack:
    """FakeRoles script for the pack: no tool round, minimal notes, the sample slot draft of the frozen evidence."""

    def __init__(self):
        self.ev = None

    def __call__(self, role, msgs, tools, schema):
        if schema is None:
            return AIMessage(content="Findings ready.")
        if schema is m.InsightNotes:
            return m.InsightNotes(summary=[m.Cited(text="see the KPI evidence", evidence_ids=["kpi.weekly"])])
        w = int(self.key.split(":")[1])
        d = sample.slot_draft(w, self.ev)
        # the ideal Writer says when the top driver is under its budget (findings v3: test week 36's planted miss)
        c1 = sample.top3(self.ev)[0][0]
        var = next(r["budget_variance_sdr"] for r in self.ev["kpi.category"]["rows"] if r["vessel_category"] == c1)
        if var is not None and var < 0:
            for b in d["bullets"]:
                if b["slide"] == 5:
                    b["text"] = b["text"].replace(", budget ", ", below its budget of ")
        return m.DeckDraft.model_validate(d)


class FakeQa:
    """FakeRoles script for the Q&A: the tool call of the item's first value element (or the obvious one), then the
    answer written from the golden elements as slots."""

    TOOL_FOR = {"q19": ("finance_kpi", {"report_week": 41, "dataset": "customers"}),
                "q21": ("invoice_audit", {"report_week": 15}),
                "r01": ("finance_kpi", {"report_week": 23, "dataset": "weekly"}),
                "r02": ("finance_kpi", {"report_week": 49, "dataset": "overdue"})}

    def __init__(self):
        self.item = None

    def _call(self):
        it = self.item
        el = next((e for e in it["elements"] if e["kind"] == "value"), None)
        if el is not None:
            args = dict(el["args"])
            if el["tool"] == "compute_toll":
                args["scnt_tons"] = 50000.0
            return el["tool"], args
        return self.TOOL_FOR.get(it["twin_of"] if "twin_of" in it else it["id"])

    def __call__(self, role, msgs, tools, schema):
        call = self._call()
        if schema is None:
            if any(x.type == "tool" for x in msgs) or call is None:
                return AIMessage(content="done")
            return AIMessage(content="", tool_calls=[{"name": call[0], "args": call[1], "id": "c1"}])
        it = self.item
        if it["expected_status"] != "answered":
            return m.QaAnswer(status=it["expected_status"], answer="Not available for this question.")
        parts = []
        for el in it["elements"]:
            if el["kind"] == "value":
                f = el["field_any"][0]
                parts.append(slots.binding("insight.c1", f, None, el["row"]))
            elif el["kind"] in ("name", "direction"):
                parts.append(el["text_any"][0])
        w = it["relevance"]["week"]
        text = " ".join([*parts, f"week {w}" if w else "", it["relevance"]["metric_any"][0]])
        return m.QaAnswer(status="answered", answer=text, evidence_ids=["insight.c1"])


# %% Phase 3 — The golden weeks with fake models (0 calls)
fp = FakePack()
reports = Path(tempfile.mkdtemp(prefix="stage1_eval_fake_"))
g_fake = wp.build_graph(roles.FakeRoles(fp), checkpointer=InMemorySaver(), reports_dir=reports)
fake_packs = []
for wk in weeks:
    key = f"{wk['dataset']}:{wk['week']}"
    fp.key, fp.ev = key, expected[key]
    st, pptx = run_pack(g_fake, key, reports)
    fake_packs.append(score_pack(key, st, pptx))
shutil.rmtree(reports, ignore_errors=True)
agg3 = aggregate(fake_packs, [])
g3 = gates({**agg3, "qa_traceability": None, "abstention_recall": 1.0})
print("Phase 3:", {k: agg3[k] for k in ("packs_published", "deck_numbers", "planted_found", "false_alarm_rate",
                                        "claim_accuracy", "traceability", "narrative_recall")}, agg3["usage"])
assert all(g3.values()), g3
assert agg3["usage"]["fresh_calls"] == 0 and agg3["packs_published"] == "11 / 11"
assert agg3["planted_found"] == f"{PLANTED} / {PLANTED}", agg3["planted_found"]
print(f"✅ Phase 3 validation passed: 11 / 11 golden weeks published with fake models; deck numbers 100 %, planted "
      f"{PLANTED} / {PLANTED} (5 in the slide-8 table + the week-36 budget miss), false alarms 0, claims 100 %, "
      "traceability 100 %; 0 model calls")


# %% Phase 4 — The finance Q&A with fake answers (0 calls)
fq = FakeQa()
g_qa_fake = qa.build_graph(roles.FakeRoles(fq))
fake_answers, skipped = [], []
for item in golden_qa:
    if any(e["kind"] == "fact" for e in item["elements"]) or item.get("must_not_regex"):
        skipped.append(item["id"])                    # document questions: live retrieval (Gemini, Jina) — Phase 5
        continue
    fq.item = item
    fake_answers.append(score_qa(item, run_qa(g_qa_fake, item)))
agg4 = aggregate([], fake_answers)
print("Phase 4:", {k: agg4[k] for k in ("qa_completeness_en", "qa_completeness_ar", "qa_relevance_en",
                                        "qa_relevance_ar", "qa_numbers", "abstention_recall", "rbac_leakage")},
      "skipped:", skipped)
bad = [(a["id"], a["status"], a["problems"][:1]) for a in fake_answers if not a["status_ok"]]
assert not bad, bad
assert (agg4["qa_completeness_en"], agg4["qa_completeness_ar"], agg4["abstention_recall"], agg4["rbac_leakage"]) \
    == (1.0, 1.0, 1.0, 0) and skipped == ["r03", "r03b", "r05", "r04"]       # the document questions (live)
print(f"✅ Phase 4 validation passed: {len(fake_answers)} questions through the Q&A graph with fake answers — "
      "completeness EN / AR 1.0, abstention 1.0, RBAC leakage 0 (the Operations employees' tool calls refused), "
      "numbers bound; 0 model calls")


# %% Phase 5 — Real models (EVAL_REAL_MODELS=1)
# ADR-052: REPEATS runs; repeat 1 on the main reply cache, repeat k on its own cache folder (a fresh set of calls),
# one shared daily quota state. A quota stop (llm.Exhausted) ends the run: its partial results go to the scratch
# folder and a re-run resumes from the caches. The results file is written only for complete runs.
results_path = EVAL / f"stage1_eval_{date.today().isoformat()}{'_' + TAG if TAG else ''}.json"


def adr052(base: dict, cand: dict) -> dict:
    """ADR-052 on two results files: no hard gate regresses (a gate the baseline already fails may stay failed or
    improve, D-P7-10), no primary's mean falls by more than 0.03; a tie keeps the simpler / cheaper setup."""
    renamed = {"planted_5_of_5": "planted_all"}          # results files before findings v3 (D-P7-1)
    base_g = {renamed.get(g, g): ok for g, ok in base["gates_all_runs"].items()}
    regressed = [g for g, ok in base_g.items() if ok and not cand["gates_all_runs"][g]]
    drops = {k: round(cand["primaries_mean"][k] - base["primaries_mean"][k], 4) for k in base["primaries_mean"]}
    keep = not regressed and all(d >= -0.03 for d in drops.values())
    better = keep and (any(d > 0.03 for d in drops.values())
                       or sum(cand["gates_all_runs"].values()) > sum(base_g.values()))
    return {"gates_regressed": regressed, "primary_change": drops, "kept": keep,
            "verdict": "keep (better)" if better else ("tie: keep the simpler / cheaper" if keep else "reject")}
if not REAL:
    print("Phase 5: skipped (set EVAL_REAL_MODELS=1 to run the real models)")
else:
    SCRATCH.mkdir(parents=True, exist_ok=True)
    quota = llm.QuotaState(llm.CACHE_DIR)
    runs, stopped = [], None
    # D-P7-11: the free tier of the main model allows 500 requests a day (Gemini's 429, 2026-10-20). A run that
    # loses its main model mid-way would mix models (the fallback answers the rest) and could not be compared
    # (ADR-052): an item finished after the main model's daily quota ran out is dropped and the run stops; a re-run
    # the next provider day resumes from the caches with the main model.
    MAIN = {c[0].provider + "/" + c[0].model for c in roles.ROLES.values()}

    def main_models_ok() -> None:
        for name in sorted(MAIN):
            prov, model = name.split("/", 1)
            quota.check(prov, model)                    # raises llm.Exhausted when the day's quota is gone
    for r in range(1, REPEATS + 1):
        cache = (llm.CACHE_DIR / f"eval_{TAG}_r{r}" if TAG else
                 llm.CACHE_DIR if r == 1 else llm.CACHE_DIR / f"eval_r{r}")
        real = roles.Roles(cache_dir=cache, quota=quota)
        reports = Path(tempfile.mkdtemp(prefix=f"stage1_eval_r{r}_"))
        g_pack = wp.build_graph(real, checkpointer=InMemorySaver(), reports_dir=reports)
        g_qa = qa.build_graph(real)
        packs, answers = [], []
        try:
            for wk in weeks:
                key = f"{wk['dataset']}:{wk['week']}"
                t0 = time.time()
                main_models_ok()
                st, pptx = run_pack(g_pack, key, reports)
                main_models_ok()                        # still available after the pack: no model mixed in it
                p = score_pack(key, st, pptx)
                p["seconds"] = round(time.time() - t0, 1)
                packs.append(p)
                if pptx:
                    shutil.copy(pptx, SCRATCH / f"r{r}_{key.replace(':', '_')}.pptx")
                print(f"  r{r} {key}: {p['status']}, findings {len(p['findings_met'])}/{len(p['findings'])}, "
                      f"planted {len(p['planted_found'])}/{len(p['planted'])}, FA {p['false_alarms']}, "
                      f"calls {p['usage']['fresh_calls']} fresh, {p['seconds']} s")
            for item in golden_qa:
                main_models_ok()
                st_q = run_qa(g_qa, item)
                main_models_ok()
                a = score_qa(item, st_q)
                answers.append(a)
                print(f"  r{r} {item['id']}: {a['status']} (expected {a['expected']})"
                      + (f", completeness {a['completeness']:.2f}" if "completeness" in a else "")
                      + (f", LEAK" if a.get("leak") else ""))
        except llm.Exhausted as exc:
            stopped = f"repeat {r}: {str(exc)[:300]}"
        finally:
            shutil.rmtree(reports, ignore_errors=True)
        run = {"repeat": r, "cache": str(cache.relative_to(ROOT)), "packs": packs, "answers": answers,
               "complete": len(packs) == len(weeks) and len(answers) == len(golden_qa)}
        if run["complete"]:
            run["aggregate"] = aggregate(packs, answers)
            run["gates"] = gates(run["aggregate"])
        runs.append(run)
        if stopped:
            break
    out = {"plan": "Agents plan §5d; ADR-052 (two runs, hard gates, primaries −0.03)", "date": date.today().isoformat(),
           "tag": TAG or "baseline",
           "inputs_sha256": sha, "models": {k: [c.name for c in v] for k, v in roles.ROLES.items()},
           "runs": runs, "stopped": stopped}
    complete = [x for x in runs if x["complete"]]
    if len(complete) == REPEATS:
        primaries = {k: round(sum(x["aggregate"][k] for x in complete) / len(complete), 4)
                     for k in ("narrative_recall", "qa_completeness_en", "qa_completeness_ar")}
        out["primaries_mean"] = primaries
        out["gates_all_runs"] = {g: all(x["gates"][g] for x in complete) for g in complete[0]["gates"]}
        out["g2_arabic_threshold"] = (primaries["qa_completeness_ar"] >= primaries["qa_completeness_en"] - 0.10)
        out["d7_met"] = all(out["gates_all_runs"].values())
        if BASELINE:
            base = json.loads((EVAL / BASELINE).read_text(encoding="utf-8"))
            out["adr052_vs"] = {"baseline": BASELINE, **adr052(base, out)}
            print("ADR-052 vs", BASELINE, ":", out["adr052_vs"])
        assert not results_path.exists(), f"{results_path.name} exists: a results file is never overwritten"
        results_path.write_text(json.dumps(out, indent=1, ensure_ascii=False, default=str), encoding="utf-8",
                                newline="\n")
        for x in complete:
            print(f"run {x['repeat']}:", json.dumps({k: v for k, v in x["aggregate"].items() if k != "style"},
                                                     ensure_ascii=False))
            print(f"run {x['repeat']} gates:", x["gates"])
        print("primaries (mean of the runs):", primaries, "| G2 AR ≥ EN − 0.10:", out["g2_arabic_threshold"])
        print(f"D7 verdict: {'MET' if out['d7_met'] else 'NOT MET'} — results file {results_path.name}")
        print("✅ Phase 5 validation passed: every run complete and scored; results file written (the baseline)")
    else:
        partial = SCRATCH / "stage1_eval_partial.json"
        partial.write_text(json.dumps(out, indent=1, ensure_ascii=False, default=str), encoding="utf-8")
        print(f"Phase 5 INCOMPLETE ({stopped}); partial results → {partial.relative_to(ROOT)}; re-run to resume "
              "from the caches")


# %% Final Checkpoint
# Narrative: pipelines/checkpoints/agents/Stage1_eval.md
print("✅ Final Checkpoint: inputs frozen, scorers proven, the harness end to end with fake models (packs + Q&A, "
      "0 calls)" + ("; real models: see Phase 5" if REAL else "; real models not run"))
