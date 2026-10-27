"""V8 — the rendered PPTX read back (Agents plan §5c, P6-7).

The deck the Finance manager approves is the file, not the draft. `validate_pptx` reopens the file and compares
every shape with the renderer's spec rebuilt from the same draft and evidence:
- every shape carries a template name (`tpl:` `bind:` `txt:` `tbl:` `chart:` `deco:`), none is missing, none is
  extra, none appears twice — a text box added by hand is caught;
- `tpl:` text = the template text exactly (the full template allowlist, by construction);
- `bind:` text = its field in its format, re-read from the evidence through the binding in the shape's name
  (independent of the spec), `txt:` = the Writer's text with its slots filled;
- table cells = the spec; a bound cell is also re-read from its field;
- chart categories, series names and values = the tool's exact values;
- `deco:` shapes hold no text.
"""
import io
from pathlib import Path

from pptx import Presentation

from pipelines.agents.deck import render, slots
from pipelines.agents.graphs import models as m

PREFIXES = ("tpl:", "bind:", "txt:", "tbl:", "chart:", "deco:")


def _text(shape) -> str:
    return "\n".join(p.text for p in shape.text_frame.paragraphs)


def read(data: bytes | Path) -> list[list[tuple[str, dict]]]:
    """Every slide as (shape name, what it shows), in drawing order."""
    prs = Presentation(io.BytesIO(data) if isinstance(data, bytes) else str(data))
    out = []
    for slide in prs.slides:
        items = []
        for sh in slide.shapes:
            if getattr(sh, "has_chart", False) and sh.has_chart:
                plot = sh.chart.plots[0]
                items.append((sh.name, {"kind": "chart", "categories": [str(c) for c in plot.categories],
                                        "series": [(s.name, list(s.values)) for s in plot.series]}))
            elif getattr(sh, "has_table", False) and sh.has_table:
                t = sh.table
                items.append((sh.name, {"kind": "table", "cells": [[t.cell(i, j).text for j in range(len(t.columns))]
                                                                   for i in range(len(t.rows))]}))
            elif sh.has_text_frame:
                items.append((sh.name, {"kind": "text", "text": _text(sh)}))
            else:
                items.append((sh.name, {"kind": "other"}))
        out.append(items)
    return out


def _num(v):
    return None if v is None else float(v)


def validate_pptx(data: bytes | Path, draft: dict, evidence: dict, readiness: dict, label: str) -> list[m.Problem]:
    problems: list[m.Problem] = []

    def bad(slide, name, message):
        problems.append(m.Problem(rule="V8", slide=slide, token=name, message=message))

    spec = render.build_spec(draft, evidence, readiness, label)
    try:
        got = read(data)
    except Exception as exc:                                    # a file that does not open is a failed deck
        return [m.Problem(rule="V8", message=f"the PPTX does not open: {type(exc).__name__}: {str(exc)[:200]}")]
    if len(got) != len(spec):
        bad(None, None, f"{len(got)} slides, the deck has {len(spec)}")
    for s, (els, items) in enumerate(zip(spec, got), start=1):
        want = {e.name: e for e in els}
        seen = set()
        for name, it in items:
            if not name.startswith(PREFIXES):
                bad(s, name, f"a shape without a template name ({it.get('text', it['kind'])[:40]!r})")
                continue
            if name in seen:
                bad(s, name, "appears twice")
                continue
            seen.add(name)
            e = want.get(name)
            if e is None:
                bad(s, name, "is not in the deck's template")
                continue
            _compare(s, e, it, evidence, bad)
        for name in want.keys() - seen:
            bad(s, name, "is missing from the file")
    return problems


def _compare(s, e, it, evidence, bad):
    if e.kind in ("text", "deco"):
        if it["kind"] != "text":
            return bad(s, e.name, f"is a {it['kind']}, the template has text")
        if it["text"] != e.text:
            return bad(s, e.name, f"shows {it['text'][:60]!r}, the deck says {e.text[:60]!r}")
        if e.name.startswith("bind:") and it["text"] != slots.bound_text(e.name[5:], evidence):
            return bad(s, e.name, f"{it['text']!r} is not its field in its format")
    elif e.kind == "table":
        if it["kind"] != "table":
            return bad(s, e.name, f"is a {it['kind']}, the template has a table")
        want = [[c.text for c in row] for row in e.cells]
        if [len(r) for r in it["cells"]] != [len(r) for r in want]:
            return bad(s, e.name, "has another shape of rows / columns")
        for i, (row, wrow) in enumerate(zip(it["cells"], e.cells)):
            for j, (txt, c) in enumerate(zip(row, wrow)):
                if txt != c.text or (c.bind and txt != slots.bound_text(c.bind, evidence)):
                    bad(s, e.name, f"cell ({i}, {j}) shows {txt[:40]!r}, the deck says {c.text[:40]!r}")
    elif e.kind == "chart":
        if it["kind"] != "chart":
            return bad(s, e.name, f"is a {it['kind']}, the template has a chart")
        c = e.chart
        if it["categories"] != [str(x) for x in c["categories"]]:
            bad(s, e.name, "its categories differ from the evidence")
        want = [(n, [_num(v) for v in vals]) for n, vals in c["series"]]
        have = [(n, [_num(v) for v in vals]) for n, vals in it["series"]]
        if have != want:
            diff = next(((n, i, a, b) for (n, va), (_, vb) in zip(have, want) for i, (a, b) in enumerate(zip(va, vb))
                         if a != b), None)
            bad(s, e.name, f"a point differs from the evidence: {diff}" if diff else "its series differ")
