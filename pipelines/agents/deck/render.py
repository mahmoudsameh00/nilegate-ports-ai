"""The NileGate weekly deck renderer (Agents plan §5c, P6-2 / P6-4 / P6-7): draft + evidence → a 10-slide PPTX.

Two steps, so the Validator can check the file against the same intent the renderer drew from:
1. `build_spec(draft, evidence, readiness, label)` — every element of every slide as data: its **name**, its kind
   and what it must show. Names carry the contract V8 checks:
   - `tpl:<key>`   fixed template text (brand, footers, labels, headers, footnotes) — the full template allowlist;
   - `bind:{…}`    one code-rendered value, the slot text of its field (a tile's number);
   - `txt:<key>`   the Writer's text with its slots filled (`txt:title`, `txt:b<i>` = bullet i of the draft);
   - `tbl:<name>`  a table: header cells are template text, data cells are labels from the evidence or bindings;
   - `chart:<name>` a native chart: categories and series = the tool's exact values (axis formats show M / k);
   - `deco:<key>`  a shape without text (backgrounds, rules).
2. `render(...)` draws the spec with python-pptx and returns the file bytes. No number is computed here: values come
   from the evidence through `slots` (the same formatting as the Writer's slots), so the deck and the draft agree.
The code-rendered parts follow deck-style-guide (slides 1, 3 tiles, every table and chart, the footnotes, slide 10).
"""
import io
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from pipelines.agents.deck import slots
from pipelines.agents.deck import style as S

SLIDES = 10
BRAND = "NileGate Weekly Finance Pack"
WARMUP_FOOTNOTE = "Receivables still building up since 1 Jan 2025"
CONTEXT_NOTE = "External context: public data and news, not NileGate figures"
UNIT_NOTE = "Amounts in SDR"
# proposed action per audit check (skill suez-toll-audit)
ACTIONS = {"duplicate_charge": "Void it; if paid, credit note / refund",
           "tonnage_mismatch": "Correct the tonnage",
           "amount_mismatch": "Corrected invoice for the difference",
           "rate_mismatch": "Correct the rate"}
RECORD_ONLY = "Record to correct (no money at stake)"
AGING = (("aging_current_sdr", "Current"), ("aging_1_30_sdr", "1–30 days"), ("aging_31_60_sdr", "31–60 days"),
         ("aging_61_90_sdr", "61–90 days"), ("aging_90_plus_sdr", "90+ days"))
TILES = (("Billed", "billed_sdr", "M", "billed_wow_pct", "WoW"),
         ("Collected", "collected_sdr", "M", "collected_wow_pct", "WoW"),
         ("AR", "ar_sdr", "M", "ar_wow_pct", "WoW"),
         ("Overdue", "overdue_sdr", "M", "overdue_wow_pct", "WoW"),
         ("DSO", "dso_days", None, "dso_wow_days", "WoW"),
         ("Budget variance", "budget_variance_sdr", "M", "budget_variance_pct", "vs budget"))


@dataclass
class Cell:
    text: str
    bind: str | None = None          # slot text when the cell is a tool number


@dataclass
class El:
    name: str
    kind: str                        # text | table | chart | deco
    box: tuple                       # left, top, width, height (EMU)
    text: str = ""
    style: dict = field(default_factory=dict)
    cells: list = None               # table: rows of Cell
    widths: tuple = ()               # table column widths (fractions)
    chart: dict = None               # {"type", "categories", "series": [(name, values)], "fmt", "highlight"}


def _in(*v):
    return tuple(Inches(x) for x in v)


def _bind(eid, f, fmt=None, row=None, direction=False) -> str:
    return slots.binding(eid, f, fmt, row, direction)


def _cell(ev, eid, f, fmt=None, row=None) -> Cell:
    b = _bind(eid, f, fmt, row)
    return Cell(slots.bound_text(b, ev), b)


# --- the spec -------------------------------------------------------------------------------------------------------

class _Spec:
    def __init__(self, draft: dict, evidence: dict, readiness: dict, label: str):
        self.d, self.ev, self.ready, self.label = draft, evidence, readiness or {}, label
        self.k = evidence["kpi.weekly"]["rows"][0]
        self.week_label = self.k["iso_week_label"]
        self.titles = {st["slide"]: slots.render_text(st["title"], evidence)[0] for st in draft["slide_titles"]}
        self.bullets = {}
        for i, b in enumerate(draft["bullets"]):
            self.bullets.setdefault(b["slide"], []).append((i, slots.render_text(b["text"], evidence)[0]))

    # common parts
    def frame(self, s: int) -> list[El]:
        y = S.FOOTER_TOP
        out = [El("txt:title", "text", (S.MARGIN, S.TITLE_TOP, S.CONTENT_W, S.TITLE_H), self.titles.get(s, ""),
                  {"size": S.TITLE_PT, "bold": True, "color": S.NAVY, "anchor": "bottom"}),
               El("deco:rule", "deco", (S.MARGIN, S.TITLE_TOP + S.TITLE_H + Inches(0.05), Inches(1.2), Inches(0.06)),
                  style={"fill": S.NILE}),
               El("tpl:footer", "text", (S.MARGIN, y, Inches(7), S.FOOTER_H),
                  f"NileGate · {BRAND.removeprefix('NileGate ')} · {self.week_label} · {UNIT_NOTE}",
                  {"size": S.FOOT_PT, "color": S.MUTED}),
               El("tpl:page", "text", (S.WIDTH - S.MARGIN - Inches(0.8), y, Inches(0.8), S.FOOTER_H), f"{s} / {SLIDES}",
                  {"size": S.FOOT_PT, "color": S.MUTED, "align": "right"}),
               self.badge(S.WIDTH - S.MARGIN - Inches(3.2), y)]
        return out

    def badge(self, x, y, w=Inches(2.2), h=S.FOOTER_H) -> El:
        approved = self.label != "Draft for approval"
        return El("tpl:label", "text", (x, y, w, h), self.label,
                  {"size": S.FOOT_PT, "bold": True, "color": S.WHITE if approved else S.INK, "align": "center",
                   "fill": S.APPROVED_BG if approved else S.LABEL_BG, "shape": "rounded", "anchor": "middle"})

    def comments(self, s: int, x, y, w, step=S.COMMENT_H) -> list[El]:
        return [El(f"txt:b{i}", "text", (x, y + n * step, w, step), text,
                   {"size": S.BODY_PT if s == 2 else S.SMALL_PT + Pt(2), "color": S.INK})
                for n, (i, text) in enumerate(self.bullets.get(s, []))]

    def footnote(self) -> list[El]:
        if not self.k.get("is_warmup"):
            return []
        return [El("tpl:footnote", "text", (S.MARGIN, S.FOOTNOTE_TOP, S.CONTENT_W, Inches(0.3)), WARMUP_FOOTNOTE,
                   {"size": S.FOOT_PT, "color": S.MUTED, "italic": True})]

    # slides
    def s1(self) -> list[El]:
        k = self.k
        return [El("deco:bg", "deco", (0, 0, S.WIDTH, S.HEIGHT), style={"fill": S.NAVY}),
                El("deco:rule", "deco", _in(0.9, 2.55, 1.6, 0.08), style={"fill": S.NILE}),
                El("tpl:brand", "text", _in(0.9, 1.4, 11.5, 1.0), BRAND,
                   {"size": S.BRAND_PT, "bold": True, "color": S.WHITE}),
                El("tpl:week", "text", _in(0.9, 2.8, 11.5, 0.6),
                   f"{k['iso_week_label']} · {k['week_start']} – {k['week_end']}",
                   {"size": S.WEEK_PT, "color": S.WHITE}),
                El("txt:title", "text", _in(0.9, 3.6, 11.5, 1.2), self.titles.get(1, ""),
                   {"size": Pt(24), "color": S.WHITE}),
                El("tpl:unit", "text", _in(0.9, 6.2, 6, 0.4), f"Finance · weekly report · {UNIT_NOTE}",
                   {"size": S.SMALL_PT, "color": S.WHITE}),
                self.badge(*_in(10.2, 6.15, 2.3, 0.45))]

    def s2(self) -> list[El]:
        out = self.frame(2)
        for n, (i, text) in enumerate(self.bullets.get(2, [])):
            y = S.BODY_TOP + Inches(0.1 + n * 0.95)
            out += [El(f"deco:mark{n}", "deco", (S.MARGIN, y + Inches(0.12), Inches(0.14), Inches(0.14)),
                       style={"fill": S.NILE}),
                    El(f"txt:b{i}", "text", (S.MARGIN + Inches(0.35), y, S.CONTENT_W - Inches(0.35), Inches(0.85)),
                       text, {"size": S.BODY_PT + Pt(2), "color": S.INK})]
        return out

    def s3(self) -> list[El]:
        out, gap, th = self.frame(3), Inches(0.25), Inches(1.45)
        tw = Emu(int((S.CONTENT_W - 2 * gap) / 3))
        for n, (label, f, fmt, ch, ch_label) in enumerate(TILES):
            x, y = S.MARGIN + (n % 3) * (tw + gap), S.BODY_TOP + (n // 3) * (th + gap)
            vb, cb = _bind("kpi.weekly", f, fmt), _bind("kpi.weekly", ch, "%2" if ch.endswith("_pct") else None)
            out += [El(f"deco:tile{n}", "deco", (x, y, tw, th), style={"fill": S.TILE_BG, "edge": S.NILE}),
                    El(f"tpl:tile{n}.label", "text", (x + Inches(0.2), y + Inches(0.1), tw, Inches(0.35)), label,
                       {"size": S.TILE_LABEL_PT, "color": S.INK_2, "bold": True}),
                    El(f"bind:{vb}", "text", (x + Inches(0.2), y + Inches(0.42), tw - Inches(0.3), Inches(0.55)),
                       slots.bound_text(vb, self.ev), {"size": S.TILE_VALUE_PT, "bold": True, "color": S.INK}),
                    El(f"bind:{cb}", "text", (x + Inches(0.2), y + Inches(1.0), Inches(1.6), Inches(0.35)),
                       slots.bound_text(cb, self.ev), {"size": S.TILE_LABEL_PT, "color": S.INK}),
                    El(f"tpl:tile{n}.change", "text", (x + Inches(1.8), y + Inches(1.0), Inches(1.4), Inches(0.35)),
                       ch_label, {"size": S.TILE_LABEL_PT, "color": S.MUTED})]
        out += self.comments(3, S.MARGIN, S.BODY_TOP + 2 * (th + gap), S.CONTENT_W)
        return out + self.footnote()

    def drivers(self):
        rows = self.ev["kpi.category"]["rows"]
        by = ("billed_wow_change_sdr" if any(r["billed_wow_change_sdr"] is not None for r in rows)
              else "budget_variance_sdr")
        return sorted(rows, key=lambda r: -abs(r[by] or 0))[:3], by

    def s4(self) -> list[El]:
        out = self.frame(4)
        trend = self.ev["kpi.trend"]["rows"]
        cats = [r["trend_week_label"] for r in trend]
        hi = next((i for i, r in enumerate(trend) if r.get("is_report_week")), len(trend) - 1)
        out.append(El("chart:trend", "chart", _in(0.6, 1.55, 6.8, 3.7), chart={
            "type": "line", "categories": cats, "fmt": S.AXIS_M, "highlight": hi,
            "series": [("Billed", [r["billed_sdr"] for r in trend]), ("Budget", [r["budget_sdr"] for r in trend])]}))
        top, by = self.drivers()
        flagged = {r["vessel_category"] for r in self.ev["audit.week"]["rows"]}
        if by == "billed_wow_change_sdr":
            head = ["Driver", "Billed", "WoW change", "WoW"]
            body = [[Cell(r["vessel_category"]), _cell(self.ev, "kpi.category", "billed_sdr", "M", r["vessel_category"]),
                     _cell(self.ev, "kpi.category", "billed_wow_change_sdr", "M", r["vessel_category"]),
                     _cell(self.ev, "kpi.category", "billed_wow_pct", "%1", r["vessel_category"])] for r in top]
        else:
            head = ["Driver (week one: vs budget)", "Billed", "vs budget", "vs budget %"]
            body = [[Cell(r["vessel_category"]), _cell(self.ev, "kpi.category", "billed_sdr", "M", r["vessel_category"]),
                     _cell(self.ev, "kpi.category", "budget_variance_sdr", "M", r["vessel_category"]),
                     _cell(self.ev, "kpi.category", "budget_variance_pct", "%1", r["vessel_category"])] for r in top]
        for row, r in zip(body, top):
            row.append(Cell("audit finding" if r["vessel_category"] in flagged else ""))
        out.append(El("tbl:drivers", "table", _in(7.6, 1.6, 5.13, 1.9), cells=[[Cell(h) for h in head + ["Audit"]]]
                      + body, widths=(0.33, 0.17, 0.19, 0.13, 0.18), style={"size": Pt(10)}))
        out += self.comments(4, S.MARGIN, Inches(5.4), S.CONTENT_W)
        return out

    def s5(self) -> list[El]:
        out = self.frame(5)
        # a bar chart draws its first category at the bottom: smallest first puts the largest on top (no reversed
        # axis, which moves the value axis to the top)
        largest_first = sorted(self.ev["kpi.category"]["rows"],
                               key=lambda r: (-(r["billed_sdr"] or 0), r["vessel_category"]))
        rows = largest_first[::-1]
        out.append(El("chart:mix", "chart", _in(0.6, 1.5, 8.3, 5.0), chart={
            "type": "bar", "categories": [r["vessel_category"] for r in rows], "fmt": S.AXIS_M,
            "series": [("Billed", [r["billed_sdr"] for r in rows]), ("Budget", [r["budget_sdr"] for r in rows])]}))
        out += self.comments(5, Inches(9.2), S.BODY_TOP + Inches(0.2), Inches(3.5), Inches(1.4))
        return out

    def s6(self) -> list[El]:
        out = self.frame(6)
        out.append(El("chart:aging", "chart", _in(0.6, 1.5, 5.4, 3.0), chart={
            "type": "column", "categories": [lbl for _, lbl in AGING], "fmt": S.AXIS_M,
            "series": [("Receivables by age", [self.k[f] for f, _ in AGING])]}))
        dso, dso_ch = _bind("kpi.weekly", "dso_days"), _bind("kpi.weekly", "dso_wow_days")
        out += [El("tpl:dso.label", "text", _in(0.6, 4.6, 1.0, 0.4), "DSO",
                   {"size": S.SMALL_PT + Pt(2), "bold": True, "color": S.INK_2}),
                El(f"bind:{dso}", "text", _in(1.5, 4.55, 1.8, 0.45), slots.bound_text(dso, self.ev),
                   {"size": Pt(20), "bold": True, "color": S.INK}),
                El(f"bind:{dso_ch}", "text", _in(3.4, 4.62, 1.6, 0.4), slots.bound_text(dso_ch, self.ev),
                   {"size": S.SMALL_PT + Pt(2), "color": S.INK}),
                El("tpl:dso.change", "text", _in(4.9, 4.62, 1.0, 0.4), "WoW", {"size": S.SMALL_PT, "color": S.MUTED})]
        od = self.ev["kpi.overdue"]["rows"][:5]
        if od:
            body = [[Cell(r["invoice_number"]), Cell(r["customer_name"]),
                     _cell(self.ev, "kpi.overdue", "amount_sdr", None, r["invoice_number"]),
                     _cell(self.ev, "kpi.overdue", "days_past_due", None, r["invoice_number"]),
                     Cell(r["aging_bucket"])] for r in od]
            out.append(El("tbl:overdue", "table", _in(6.3, 1.55, 6.45, 2.6), cells=[[Cell(h) for h in (
                "Largest overdue", "Customer", "Amount", "Past due", "Bucket")]] + body,
                widths=(0.2, 0.3, 0.2, 0.14, 0.16)))
        else:
            out.append(El("tpl:overdue.none", "text", _in(6.3, 1.55, 6.45, 0.5), "No overdue invoice this week",
                          {"size": S.BODY_PT, "color": S.INK_2}))
        out += self.comments(6, S.MARGIN, Inches(5.25), S.CONTENT_W)
        return out + self.footnote()

    def s7(self) -> list[El]:
        out = self.frame(7)
        cu = self.ev["kpi.customers"]["rows"][:10]
        body = [[_cell(self.ev, "kpi.customers", "rank_13w", None, r["customer_name"]),
                 _cell(self.ev, "kpi.customers", "rank_13w_prev", None, r["customer_name"]), Cell(r["customer_name"]),
                 _cell(self.ev, "kpi.customers", "billed_13w_sdr", "M", r["customer_name"]),
                 _cell(self.ev, "kpi.customers", "share_13w_pct", "%1", r["customer_name"])] for r in cu]
        out.append(El("tbl:customers", "table", _in(0.6, 1.5, 8.0, 4.6), cells=[[Cell(h) for h in (
            "Rank", "Previous", "Customer", "Billed, 13 weeks", "Share")]] + body,
            widths=(0.09, 0.11, 0.48, 0.18, 0.14)))
        share = _bind("kpi.weekly", "top10_share_13w_pct", "%1")
        out += [El("deco:tile", "deco", _in(9.0, 1.5, 3.7, 1.4), style={"fill": S.TILE_BG, "edge": S.NILE}),
                El("tpl:share.label", "text", _in(9.2, 1.6, 3.4, 0.4), "Top-10 share, 13 weeks",
                   {"size": S.TILE_LABEL_PT, "bold": True, "color": S.INK_2}),
                El(f"bind:{share}", "text", _in(9.2, 2.0, 3.4, 0.6), slots.bound_text(share, self.ev),
                   {"size": S.TILE_VALUE_PT, "bold": True, "color": S.INK})]
        out += self.comments(7, Inches(9.0), Inches(3.2), Inches(3.7), Inches(1.0))
        return out

    def s8(self) -> list[El]:
        out = self.frame(8)
        for n, (label, f, fmt) in enumerate((("Invoices checked", "invoices_checked", None),
                                             ("Flagged", "invoices_flagged", None),
                                             ("Amount at stake", "amount_at_stake_sdr", None))):
            x = S.MARGIN + n * Inches(4.1)
            b = _bind("audit.week", f, fmt)
            out += [El(f"deco:tile{n}", "deco", (x, S.BODY_TOP, Inches(3.8), Inches(1.1)),
                       style={"fill": S.TILE_BG, "edge": S.NILE}),
                    El(f"tpl:tile{n}.label", "text", (x + Inches(0.2), S.BODY_TOP + Inches(0.08), Inches(3.4),
                                                      Inches(0.35)), label,
                       {"size": S.TILE_LABEL_PT, "bold": True, "color": S.INK_2}),
                    El(f"bind:{b}", "text", (x + Inches(0.2), S.BODY_TOP + Inches(0.42), Inches(3.4), Inches(0.55)),
                       slots.bound_text(b, self.ev), {"size": Pt(22), "bold": True, "color": S.INK})]
        rows = self.ev["audit.week"]["rows"][:8]
        if rows:
            body = []
            for r in rows:
                key = f"{r['invoice_number']}/{r['check_code']}"
                action = (RECORD_ONLY if r["check_code"] == "tonnage_mismatch" and not r["difference_sdr"]
                          else ACTIONS.get(r["check_code"], "Manual review"))
                body.append([Cell(r["invoice_number"]), Cell(r["vessel_category"]), Cell(r["check_code"]),
                             _cell(self.ev, "audit.week", "difference_sdr", None, key), Cell(action)])
            out.append(El("tbl:audit", "table", _in(0.6, 2.95, 12.13, 2.2), cells=[[Cell(h) for h in (
                "Invoice", "Vessel category", "Check", "Difference", "Proposed action")]] + body,
                widths=(0.13, 0.27, 0.15, 0.15, 0.30)))
        else:
            out.append(El("tpl:audit.none", "text", _in(0.6, 3.0, 12.13, 0.5), "No invoice flagged this week",
                          {"size": S.BODY_PT, "color": S.INK_2}))
        out += self.comments(8, S.MARGIN, Inches(5.35), S.CONTENT_W)
        return out

    def s9(self) -> list[El]:
        out = self.frame(9)
        out += self.comments(9, S.MARGIN, S.BODY_TOP + Inches(0.1), S.CONTENT_W, Inches(0.8))
        out.append(El("tpl:context.note", "text", (S.MARGIN, S.FOOTNOTE_TOP, S.CONTENT_W, Inches(0.3)), CONTEXT_NOTE,
                      {"size": S.FOOT_PT, "color": S.MUTED, "italic": True}))
        return out

    def s10(self) -> list[El]:
        out = self.frame(10)
        body = []
        flagged = False
        for eid, res in sorted(self.ev.items()):
            src = res.get("source") or {}
            flagged |= any(r.get("suspected_injection") for r in res.get("rows") or [])
            body.append([Cell(eid), Cell(res.get("tool") or ""), Cell(src.get("name") or ""),
                         Cell(src.get("filter") or ""), Cell(res.get("status") or ""),
                         Cell(src.get("retrieved_at") or "")])
        row_h = 0.43                                         # a 10 pt row as PowerPoint draws it
        out.append(El("tbl:sources", "table", _in(0.6, 1.5, 12.13, row_h * (len(body) + 1)), cells=[[Cell(h) for h in (
            "Evidence", "Tool", "Source", "Filter", "Status", "Retrieved")]] + body,
            widths=(0.17, 0.12, 0.2, 0.27, 0.1, 0.14), style={"size": Pt(10)}))
        facts = self.ready.get("facts") or {}
        missing = [eid for eid, res in sorted(self.ev.items()) if res.get("status") != "ok"]
        lines = [f"Data as of {facts.get('week_end')}; payments loaded through {facts.get('payments_through')}",
                 "Not checked: " + (", ".join(missing) if missing else "nothing — every source answered")]
        if flagged:
            lines.append("A retrieved passage was flagged as a possible injection and used as data only")
        y = Inches(1.5) + Inches(row_h) * (len(body) + 1) + Inches(0.25)
        for n, line in enumerate(lines):
            out.append(El(f"tpl:fresh{n}", "text", (S.MARGIN, y + n * Inches(0.38), S.CONTENT_W, Inches(0.35)), line,
                          {"size": S.SMALL_PT, "color": S.INK_2}))
        return out


def build_spec(draft: dict, evidence: dict, readiness: dict, label: str) -> list[list[El]]:
    sp = _Spec(draft, evidence, readiness, label)
    return [getattr(sp, f"s{s}")() for s in range(1, SLIDES + 1)]


# --- drawing --------------------------------------------------------------------------------------------------------

def _text(slide, e: El):
    st = e.style
    if st.get("fill") is not None:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if st.get("shape") == "rounded"
                                     else MSO_SHAPE.RECTANGLE, *e.box)
        shp.fill.solid()
        shp.fill.fore_color.rgb = st["fill"]
        shp.line.fill.background()
        shp.shadow.inherit = False
    else:
        shp = slide.shapes.add_textbox(*e.box)
    shp.name = e.name
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.vertical_anchor = {"bottom": MSO_ANCHOR.BOTTOM, "middle": MSO_ANCHOR.MIDDLE}.get(st.get("anchor"),
                                                                                    MSO_ANCHOR.TOP)
    p = tf.paragraphs[0]
    p.alignment = {"right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}.get(st.get("align"), PP_ALIGN.LEFT)
    r = p.add_run()
    r.text = e.text
    f = r.font
    f.name, f.size, f.bold, f.italic = S.FONT, st.get("size", S.BODY_PT), st.get("bold", False), st.get("italic",
                                                                                                        False)
    f.color.rgb = st.get("color", S.INK)


def _deco(slide, e: El):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *e.box)
    shp.name = e.name
    shp.fill.solid()
    shp.fill.fore_color.rgb = e.style["fill"]
    if e.style.get("edge") is not None:
        shp.line.color.rgb = e.style["edge"]
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False


def _table(slide, e: El):
    rows, cols = len(e.cells), len(e.cells[0])
    gf = slide.shapes.add_table(rows, cols, e.box[0], e.box[1], e.box[2], Emu(int(Inches(0.32) * rows)))
    gf.name = e.name
    t = gf.table
    total = e.box[2]
    for j, frac in enumerate(e.widths or [1 / cols] * cols):
        t.columns[j].width = Emu(int(total * frac))
    size = e.style.get("size", Pt(11))
    for i, row in enumerate(e.cells):
        for j, c in enumerate(row):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = S.NAVY if i == 0 else (S.WHITE if i % 2 else S.TILE_BG)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if (c.bind and i) else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = c.text
            r.font.name, r.font.size, r.font.bold = S.FONT, size, i == 0
            r.font.color.rgb = S.WHITE if i == 0 else S.INK


def _chart(slide, e: El):
    c = e.chart
    data = CategoryChartData(number_format="General")
    data.categories = c["categories"]
    for name, values in c["series"]:
        data.add_series(name, values)
    kind = {"line": XL_CHART_TYPE.LINE_MARKERS, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED}[c["type"]]
    gf = slide.shapes.add_chart(kind, *e.box, data)
    gf.name = e.name
    ch = gf.chart
    ch.font.name, ch.font.size = S.FONT, Pt(11)
    ch.font.color.rgb = S.INK_2
    ch.has_legend = len(c["series"]) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
    va = ch.value_axis
    va.tick_labels.number_format = c["fmt"]
    va.tick_labels.number_format_is_linked = False
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = S.GRID
    va.format.line.fill.background()
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(10)
    ca.format.line.color.rgb = S.GRID
    plot = ch.plots[0]
    colors = [S.SERIES_1, S.COMPARISON]
    if c["type"] == "line":
        for n, s in enumerate(plot.series):
            s.smooth = False
            s.format.line.color.rgb = colors[n]
            s.format.line.width = Pt(2.25)
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = 5
            s.marker.format.fill.solid()
            s.marker.format.fill.fore_color.rgb = colors[n]
            s.marker.format.line.color.rgb = colors[n]
            if n:
                s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        hi = c.get("highlight")
        if hi is not None and plot.series:
            pt = plot.series[0].points[hi]
            pt.marker.style = XL_MARKER_STYLE.CIRCLE
            pt.marker.size = 10
            pt.marker.format.fill.solid()
            pt.marker.format.fill.fore_color.rgb = S.HIGHLIGHT
            pt.marker.format.line.color.rgb = S.HIGHLIGHT
    else:
        plot.gap_width = 60
        if len(plot.series) > 1:
            plot.overlap = -10
        for n, s in enumerate(plot.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = colors[n]


def render(draft: dict, evidence: dict, readiness: dict, label: str) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = S.WIDTH, S.HEIGHT
    blank = prs.slide_layouts[6]
    for els in build_spec(draft, evidence, readiness, label):
        slide = prs.slides.add_slide(blank)
        for e in els:
            {"text": _text, "deco": _deco, "table": _table, "chart": _chart}[e.kind](slide, e)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
